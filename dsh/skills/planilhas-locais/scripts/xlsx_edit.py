#!/usr/bin/env python3
"""Apply explicit, narrow XLSX edits with openpyxl and reopen validation."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any


class WorkbookError(ValueError):
    pass


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def require_sheet(book: Any, name: Any) -> Any:
    if not isinstance(name, str) or name not in book.sheetnames:
        raise WorkbookError(f"sheet not found: {name!r}")
    return book[name]


def cells_in_range(sheet: Any, ref: str) -> tuple[list[Any], int, int]:
    try:
        rows = sheet[ref]
    except (KeyError, TypeError, ValueError) as exc:
        raise WorkbookError(f"invalid range {ref!r}") from exc
    if hasattr(rows, "coordinate"):
        return [rows], 1, 1
    if not isinstance(rows, tuple):
        rows = (rows,)
    if rows and not isinstance(rows[0], tuple):
        rows = (rows,)
    height = len(rows)
    width = len(rows[0]) if rows else 0
    if any(len(row) != width for row in rows):
        raise WorkbookError(f"invalid range {ref!r}")
    return [cell for row in rows for cell in row], height, width


def load_ops(path: Path | None) -> list[dict[str, Any]]:
    if not path:
        return []
    payload = json.loads(path.read_text(encoding="utf-8"))
    operations = payload.get("operations") if isinstance(payload, dict) else payload
    if not isinstance(operations, list):
        raise WorkbookError("operations file must be a list or an object with 'operations'")
    return operations


def snapshot(book: Any) -> dict[tuple[str, str], Any]:
    result: dict[tuple[str, str], Any] = {}
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    result[(sheet.title, cell.coordinate)] = cell.value
    return result


def apply(book: Any, operations: list[dict[str, Any]]) -> set[tuple[str, str]]:
    touched: set[tuple[str, str]] = set()
    for number, operation in enumerate(operations, start=1):
        if not isinstance(operation, dict) or not isinstance(operation.get("op"), str):
            raise WorkbookError(f"operation {number} must contain an 'op' string")
        name = operation["op"]
        if name == "set_cells":
            entries = operation.get("cells")
            if not isinstance(entries, list) or not entries:
                raise WorkbookError("set_cells requires a non-empty cells list")
            for entry in entries:
                if not isinstance(entry, dict) or not isinstance(entry.get("cell"), str):
                    raise WorkbookError("each set_cells entry needs cell and value")
                sheet = require_sheet(book, entry.get("sheet"))
                cell = sheet[entry["cell"]]
                cell.value = entry.get("value")
                touched.add((sheet.title, cell.coordinate))
        elif name == "set_range":
            sheet = require_sheet(book, operation.get("sheet"))
            ref, values = operation.get("range"), operation.get("values")
            if not isinstance(ref, str) or not isinstance(values, list) or not values or any(not isinstance(row, list) for row in values):
                raise WorkbookError("set_range requires range and a rectangular values matrix")
            cells, range_height, range_width = cells_in_range(sheet, ref)
            width = len(values[0])
            if width == 0 or any(len(row) != width for row in values) or len(values) != range_height or width != range_width:
                raise WorkbookError("set_range values dimensions do not match range")
            for cell, value in zip(cells, [value for row in values for value in row]):
                cell.value = value
                touched.add((sheet.title, cell.coordinate))
        elif name == "add_table":
            try:
                from openpyxl.worksheet.table import Table, TableStyleInfo
            except ImportError as exc:
                raise WorkbookError("openpyxl table support is unavailable") from exc
            sheet = require_sheet(book, operation.get("sheet"))
            ref, table_name = operation.get("ref"), operation.get("name")
            if not isinstance(ref, str) or not isinstance(table_name, str) or not table_name:
                raise WorkbookError("add_table requires ref and name")
            if table_name in {table.name for ws in book.worksheets for table in ws.tables.values()}:
                raise WorkbookError(f"table name already exists: {table_name}")
            table = Table(displayName=table_name, ref=ref)
            style_name = operation.get("style", "TableStyleMedium2")
            table.tableStyleInfo = TableStyleInfo(name=style_name, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            sheet.add_table(table)
        elif name == "create_sheet":
            sheet_name = operation.get("sheet")
            if not isinstance(sheet_name, str) or not sheet_name:
                raise WorkbookError("create_sheet requires a non-empty sheet")
            if sheet_name in book.sheetnames:
                raise WorkbookError(f"sheet already exists: {sheet_name}")
            index = operation.get("index")
            if index is not None and (not isinstance(index, int) or index < 0):
                raise WorkbookError("create_sheet index must be a non-negative integer")
            book.create_sheet(sheet_name, index=index)
        else:
            raise WorkbookError(f"unsupported operation: {name}")
    return touched


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--operations", type=Path)
    args = parser.parse_args()
    try:
        if args.input.resolve() == args.output.resolve() or (args.output.exists() and os.path.samefile(args.input, args.output)):
            raise WorkbookError("refusing to overwrite the input; choose a separate --output")
        try:
            import openpyxl
        except ImportError as exc:
            raise WorkbookError("openpyxl is required; use the DSH Python runtime with its bundled dependencies") from exc
        keep_vba = args.input.suffix.lower() == ".xlsm"
        book = openpyxl.load_workbook(args.input, data_only=False, keep_vba=keep_vba)
        original_sheets = list(book.sheetnames)
        before = snapshot(book)
        operations = load_ops(args.operations)
        touched = apply(book, operations)
        book.calculation.fullCalcOnLoad = True
        book.calculation.forceFullCalc = True
        args.output.parent.mkdir(parents=True, exist_ok=True)
        book.save(args.output)
        reopened = openpyxl.load_workbook(args.output, data_only=False, keep_vba=keep_vba)
        for key, value in before.items():
            if key not in touched and reopened[key[0]][key[1]].value != value:
                raise WorkbookError(f"untouched cell changed after save: {key[0]}!{key[1]}")
        if [*reopened.sheetnames] != original_sheets and not any(op.get("op") == "create_sheet" for op in operations):
            raise WorkbookError("sheet order/names changed unexpectedly")
        result = {
            "status": "ok", "input": str(args.input.resolve()), "output": str(args.output.resolve()),
            "input_sha256": digest(args.input), "output_sha256": digest(args.output),
            "sheets": reopened.sheetnames, "operations": [op.get("op") for op in operations],
            "touched_cells": len(touched),
            "formula_recalculation": "not performed; workbook requests recalculation when opened by a capable engine",
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (OSError, WorkbookError, json.JSONDecodeError, KeyError) as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
