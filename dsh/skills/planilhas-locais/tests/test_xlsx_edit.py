import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


SCRIPT = Path(__file__).parents[1] / "scripts" / "xlsx_edit.py"


class XlsxEditTest(unittest.TestCase):
    def test_range_edit_table_and_untouched_formula_survive_reopen(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "entrada.xlsx"
            book = Workbook()
            data = book.active
            data.title = "Dados"
            data.append(["id", "amount"])
            data.append(["001", 10])
            data.append(["002", 20])
            data["D2"] = "=SUM(B2:B3)"
            summary = book.create_sheet("Resumo")
            summary["A1"] = "mantido"
            book.save(source)
            original_formula = "=SUM(B2:B3)"
            ops = root / "ops.json"
            ops.write_text(json.dumps({"operations": [
                {"op": "set_range", "sheet": "Dados", "range": "B2:B3", "values": [[11], [22]]},
                {"op": "add_table", "sheet": "Dados", "ref": "A1:B3", "name": "TabelaDados"},
            ]}), encoding="utf-8")
            output = root / "saida.xlsx"
            result = subprocess.run([sys.executable, str(SCRIPT), "--input", str(source), "--output", str(output), "--operations", str(ops)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            reopened = load_workbook(output, data_only=False)
            self.assertEqual(reopened["Dados"]["D2"].value, original_formula)
            self.assertEqual(reopened["Resumo"]["A1"].value, "mantido")
            self.assertEqual(reopened["Dados"]["B2"].value, 11)
            self.assertIn("TabelaDados", reopened["Dados"].tables)
            self.assertEqual(load_workbook(source, data_only=False)["Dados"]["B2"].value, 10)

    def test_range_shape_mismatch_is_rejected(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "entrada.xlsx"
            book = Workbook()
            book.active.title = "Dados"
            book.save(source)
            ops = root / "ops.json"
            ops.write_text(json.dumps({"operations": [{"op": "set_range", "sheet": "Dados", "range": "A1:B2", "values": [[1, 2, 3]]}]}), encoding="utf-8")
            result = subprocess.run([sys.executable, str(SCRIPT), "--input", str(source), "--output", str(root / "saida.xlsx"), "--operations", str(ops)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 2)
            self.assertIn("dimensions", result.stderr)


if __name__ == "__main__":
    unittest.main()
